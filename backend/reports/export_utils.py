from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from vehicles.pdf_utils import register_fonts


def format_money(value):
    value = value or Decimal("0.00")

    return f"{value:.2f} zł"


def format_date(value):
    if not value:
        return "-"

    return value.strftime("%d.%m.%Y")


def calculate_service_summary(records):
    records = list(records)

    total_cost = sum(
        (
            record.cost or Decimal("0.00")
            for record in records
        ),
        Decimal("0.00"),
    )

    record_count = len(records)

    average_cost = (
        total_cost / record_count
        if record_count
        else Decimal("0.00")
    )

    most_expensive = None

    if records:
        most_expensive = max(
            records,
            key=lambda record: (
                record.cost or Decimal("0.00")
            ),
        )

    return {
        "records": records,
        "total_cost": total_cost,
        "record_count": record_count,
        "average_cost": average_cost,
        "most_expensive": most_expensive,
    }


def create_service_costs_workbook(
    service_records,
    filters,
):
    summary = calculate_service_summary(
        service_records
    )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Podsumowanie"

    summary_sheet.append(
        [
            "Raport kosztów serwisowych",
            "",
        ]
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    summary_sheet.append(
        [
            "Data od",
            filters.get("date_from") or "Wszystkie",
        ]
    )

    summary_sheet.append(
        [
            "Data do",
            filters.get("date_to") or "Wszystkie",
        ]
    )

    summary_sheet.append(
        [
            "Spółka",
            filters.get("company_name") or "Wszystkie",
        ]
    )

    summary_sheet.append(
        [
            "Oddział",
            filters.get("branch") or "Wszystkie",
        ]
    )

    summary_sheet.append(
        [
            "Pojazd",
            filters.get("vehicle_name") or "Wszystkie",
        ]
    )

    summary_sheet.append([])

    summary_sheet.append(
        [
            "Liczba wpisów serwisowych",
            summary["record_count"],
        ]
    )

    summary_sheet.append(
        [
            "Łączny koszt",
            float(summary["total_cost"]),
        ]
    )

    summary_sheet.append(
        [
            "Średni koszt",
            float(summary["average_cost"]),
        ]
    )

    if summary["most_expensive"]:
        most_expensive = (
            summary["most_expensive"]
        )

        summary_sheet.append(
            [
                "Najdroższy wpis",
                (
                    f"{most_expensive.title} - "
                    f"{most_expensive.cost:.2f} zł"
                ),
            ]
        )

    summary_sheet.column_dimensions["A"].width = 32
    summary_sheet.column_dimensions["B"].width = 45

    details_sheet = workbook.create_sheet(
        "Historia serwisowa"
    )

    headers = [
        "Data",
        "Pojazd",
        "Numer rejestracyjny",
        "Spółka",
        "Oddział",
        "Tytuł",
        "Opis",
        "Przebieg [km]",
        "Koszt [zł]",
    ]

    details_sheet.append(headers)

    for cell in details_sheet[1]:
        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for record in summary["records"]:
        vehicle = record.vehicle

        details_sheet.append(
            [
                record.service_date,
                f"{vehicle.brand} {vehicle.model}",
                vehicle.registration_number,
                (
                    vehicle.company.name
                    if vehicle.company
                    else ""
                ),
                vehicle.branch or "",
                record.title,
                record.description or "",
                record.mileage,
                float(
                    record.cost
                    or Decimal("0.00")
                ),
            ]
        )

    widths = {
        "A": 15,
        "B": 28,
        "C": 20,
        "D": 28,
        "E": 22,
        "F": 30,
        "G": 45,
        "H": 18,
        "I": 18,
    }

    for column, width in widths.items():
        details_sheet.column_dimensions[
            column
        ].width = width

    details_sheet.freeze_panes = "A2"

    if details_sheet.max_row >= 1:
        details_sheet.auto_filter.ref = (
            details_sheet.dimensions
        )

    return workbook


def create_service_costs_pdf(
    service_records,
    filters,
):
    register_fonts()

    summary = calculate_service_summary(
        service_records
    )

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Flotivo - raport kosztów serwisowych",
        author="Flotivo",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName="DejaVuSans-Bold",
        fontSize=19,
        leading=24,
        textColor=colors.HexColor("#1A1A1A"),
        spaceAfter=4 * mm,
    )

    normal_style = ParagraphStyle(
        "ReportNormal",
        parent=styles["Normal"],
        fontName="DejaVuSans",
        fontSize=8,
        leading=11,
    )

    section_style = ParagraphStyle(
        "ReportSection",
        parent=styles["Heading2"],
        fontName="DejaVuSans-Bold",
        fontSize=12,
        leading=16,
        spaceBefore=5 * mm,
        spaceAfter=3 * mm,
    )

    story = []

    story.append(
        Paragraph(
            "FLOTIVO - RAPORT KOSZTÓW SERWISOWYCH",
            title_style,
        )
    )

    filters_data = [
        [
            "Zakres dat",
            (
                f"{filters.get('date_from') or 'bez ograniczenia'}"
                f" - "
                f"{filters.get('date_to') or 'bez ograniczenia'}"
            ),
            "Spółka",
            filters.get("company_name") or "Wszystkie",
        ],
        [
            "Oddział",
            filters.get("branch") or "Wszystkie",
            "Pojazd",
            filters.get("vehicle_name") or "Wszystkie",
        ],
    ]

    filters_table = Table(
        filters_data,
        colWidths=[
            30 * mm,
            70 * mm,
            30 * mm,
            105 * mm,
        ],
    )

    filters_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "DejaVuSans",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "DejaVuSans-Bold",
                ),
                (
                    "FONTNAME",
                    (2, 0),
                    (2, -1),
                    "DejaVuSans-Bold",
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#EEE8E2"),
                ),
                (
                    "BACKGROUND",
                    (2, 0),
                    (2, -1),
                    colors.HexColor("#EEE8E2"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor("#D8D1CA"),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(filters_table)

    story.append(
        Paragraph(
            "Podsumowanie",
            section_style,
        )
    )

    most_expensive_text = "Brak"

    if summary["most_expensive"]:
        most_expensive = (
            summary["most_expensive"]
        )

        most_expensive_text = (
            f"{most_expensive.title} "
            f"({format_money(most_expensive.cost)})"
        )

    summary_data = [
        [
            "Liczba wpisów",
            str(summary["record_count"]),
            "Łączny koszt",
            format_money(
                summary["total_cost"]
            ),
        ],
        [
            "Średni koszt",
            format_money(
                summary["average_cost"]
            ),
            "Najdroższy wpis",
            most_expensive_text,
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            35 * mm,
            55 * mm,
            35 * mm,
            110 * mm,
        ],
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "DejaVuSans",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "DejaVuSans-Bold",
                ),
                (
                    "FONTNAME",
                    (2, 0),
                    (2, -1),
                    "DejaVuSans-Bold",
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#EEE8E2"),
                ),
                (
                    "BACKGROUND",
                    (2, 0),
                    (2, -1),
                    colors.HexColor("#EEE8E2"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor("#D8D1CA"),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(summary_table)

    story.append(
        Paragraph(
            "Historia serwisowa",
            section_style,
        )
    )

    if summary["records"]:
        rows = [
            [
                "Data",
                "Pojazd",
                "Rejestracja",
                "Spółka",
                "Oddział",
                "Naprawa / serwis",
                "Przebieg",
                "Koszt",
            ]
        ]

        for record in summary["records"]:
            vehicle = record.vehicle

            rows.append(
                [
                    format_date(
                        record.service_date
                    ),
                    Paragraph(
                        (
                            f"{vehicle.brand} "
                            f"{vehicle.model}"
                        ),
                        normal_style,
                    ),
                    vehicle.registration_number,
                    Paragraph(
                        (
                            vehicle.company.name
                            if vehicle.company
                            else "-"
                        ),
                        normal_style,
                    ),
                    Paragraph(
                        vehicle.branch or "-",
                        normal_style,
                    ),
                    Paragraph(
                        record.title,
                        normal_style,
                    ),
                    (
                        f"{record.mileage} km"
                        if record.mileage is not None
                        else "-"
                    ),
                    format_money(
                        record.cost
                    ),
                ]
            )

        table = Table(
            rows,
            repeatRows=1,
            colWidths=[
                22 * mm,
                39 * mm,
                30 * mm,
                43 * mm,
                34 * mm,
                54 * mm,
                27 * mm,
                28 * mm,
            ],
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, -1),
                        "DejaVuSans",
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "DejaVuSans-Bold",
                    ),
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor("#EEE8E2"),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.4,
                        colors.HexColor("#D8D1CA"),
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        story.append(table)

    else:
        story.append(
            Paragraph(
                "Brak wpisów serwisowych spełniających wybrane kryteria.",
                normal_style,
            )
        )

    story.append(
        Spacer(
            1,
            8 * mm,
        )
    )

    generated_at = timezone.localtime(
        timezone.now()
    )

    story.append(
        Paragraph(
            (
                "Raport wygenerowano w systemie Flotivo: "
                f"{generated_at.strftime('%d.%m.%Y %H:%M')}"
            ),
            normal_style,
        )
    )

    document.build(story)

    output.seek(0)

    return output