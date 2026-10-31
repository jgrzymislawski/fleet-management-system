from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from vehicles.pdf_utils import register_fonts


def format_date(value):
    if not value:
        return "-"

    return value.strftime("%d.%m.%Y")


def get_category_name(category):
    if category == "vehicle":
        return "Pojazd"

    if category == "driver":
        return "Kierowca"

    return category


def calculate_alert_summary(alerts):
    expired = 0
    within_7_days = 0
    within_30_days = 0

    for alert in alerts:
        days_left = alert["days_left"]

        if days_left < 0:
            expired += 1

        elif days_left <= 7:
            within_7_days += 1

        elif days_left <= 30:
            within_30_days += 1

    return {
        "total": len(alerts),
        "expired": expired,
        "within_7_days": within_7_days,
        "within_30_days": within_30_days,
    }


def create_alerts_report_workbook(
    alerts,
    filters,
):
    summary = calculate_alert_summary(
        alerts
    )

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Podsumowanie"

    summary_sheet.append(
        [
            "Raport terminów i alertów",
            "",
        ]
    )

    summary_sheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    summary_sheet.append([])

    summary_sheet.append(
        [
            "Kategoria",
            filters.get("category_name")
            or "Wszystkie",
        ]
    )

    summary_sheet.append(
        [
            "Pilność",
            filters.get("status_name")
            or "Wszystkie",
        ]
    )

    summary_sheet.append(
        [
            "Spółka",
            filters.get("company_name")
            or "Wszystkie",
        ]
    )

    summary_sheet.append(
        [
            "Oddział",
            filters.get("branch")
            or "Wszystkie",
        ]
    )

    summary_sheet.append(
        [
            "Pojazd",
            filters.get("vehicle_name")
            or "Wszystkie",
        ]
    )

    summary_sheet.append([])

    summary_sheet.append(
        [
            "Wszystkie alerty",
            summary["total"],
        ]
    )

    summary_sheet.append(
        [
            "Po terminie",
            summary["expired"],
        ]
    )

    summary_sheet.append(
        [
            "Wygasa w ciągu 7 dni",
            summary["within_7_days"],
        ]
    )

    summary_sheet.append(
        [
            "Wygasa za 8-30 dni",
            summary["within_30_days"],
        ]
    )

    summary_sheet.column_dimensions[
        "A"
    ].width = 32

    summary_sheet.column_dimensions[
        "B"
    ].width = 45

    details_sheet = workbook.create_sheet(
        "Alerty"
    )

    headers = [
        "Kategoria",
        "Obiekt",
        "Dodatkowe informacje",
        "Typ alertu",
        "Termin",
        "Pozostało / po terminie",
        "Status",
        "Spółka",
        "Oddział",
    ]

    details_sheet.append(
        headers
    )

    for cell in details_sheet[1]:
        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for alert in alerts:
        details_sheet.append(
            [
                get_category_name(
                    alert["category"]
                ),
                alert["name"],
                alert["secondary"],
                alert["type"],
                alert["date"],
                alert["days_text"],
                alert["status"],
                alert.get(
                    "company_name"
                ) or "",
                alert.get(
                    "branch"
                ) or "",
            ]
        )

    widths = {
        "A": 16,
        "B": 28,
        "C": 24,
        "D": 30,
        "E": 16,
        "F": 24,
        "G": 22,
        "H": 28,
        "I": 22,
    }

    for (
        column,
        width,
    ) in widths.items():
        details_sheet.column_dimensions[
            column
        ].width = width

    details_sheet.freeze_panes = (
        "A2"
    )

    if details_sheet.max_row >= 1:
        details_sheet.auto_filter.ref = (
            details_sheet.dimensions
        )

    return workbook


def create_alerts_report_pdf(
    alerts,
    filters,
):
    register_fonts()

    summary = calculate_alert_summary(
        alerts
    )

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=(
            "Flotivo - raport terminów i alertów"
        ),
        author="Flotivo",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "AlertsReportTitle",
        parent=styles["Title"],
        fontName="DejaVuSans-Bold",
        fontSize=19,
        leading=24,
        textColor=colors.HexColor(
            "#1A1A1A"
        ),
        spaceAfter=4 * mm,
    )

    section_style = ParagraphStyle(
        "AlertsReportSection",
        parent=styles["Heading2"],
        fontName="DejaVuSans-Bold",
        fontSize=12,
        leading=16,
        spaceBefore=5 * mm,
        spaceAfter=3 * mm,
    )

    normal_style = ParagraphStyle(
        "AlertsReportNormal",
        parent=styles["Normal"],
        fontName="DejaVuSans",
        fontSize=7,
        leading=10,
        textColor=colors.HexColor(
            "#1A1A1A"
        ),
    )

    story = []

    story.append(
        Paragraph(
            (
                "FLOTIVO - RAPORT "
                "TERMINÓW I ALERTÓW"
            ),
            title_style,
        )
    )

    filters_data = [
        [
            "Kategoria",
            filters.get(
                "category_name"
            )
            or "Wszystkie",
            "Pilność",
            filters.get(
                "status_name"
            )
            or "Wszystkie",
        ],
        [
            "Spółka",
            filters.get(
                "company_name"
            )
            or "Wszystkie",
            "Oddział",
            filters.get(
                "branch"
            )
            or "Wszystkie",
        ],
        [
            "Pojazd",
            filters.get(
                "vehicle_name"
            )
            or "Wszystkie",
            "",
            "",
        ],
    ]

    filters_table = Table(
        filters_data,
        colWidths=[
            30 * mm,
            75 * mm,
            30 * mm,
            100 * mm,
        ],
    )

    filters_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "DejaVuSans",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "DejaVuSans-Bold",
                ),
                (
                    "FONTNAME",
                    (2, 0),
                    (2, -1),
                    "DejaVuSans-Bold",
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#EEE8E2"
                    ),
                ),
                (
                    "BACKGROUND",
                    (2, 0),
                    (2, -1),
                    colors.HexColor(
                        "#EEE8E2"
                    ),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor(
                        "#D8D1CA"
                    ),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(
        filters_table
    )

    story.append(
        Paragraph(
            "Podsumowanie",
            section_style,
        )
    )

    summary_data = [
        [
            "Wszystkie",
            str(
                summary["total"]
            ),
            "Po terminie",
            str(
                summary["expired"]
            ),
        ],
        [
            "Do 7 dni",
            str(
                summary[
                    "within_7_days"
                ]
            ),
            "8-30 dni",
            str(
                summary[
                    "within_30_days"
                ]
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            42 * mm,
            35 * mm,
            42 * mm,
            35 * mm,
        ],
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "DejaVuSans",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "DejaVuSans-Bold",
                ),
                (
                    "FONTNAME",
                    (2, 0),
                    (2, -1),
                    "DejaVuSans-Bold",
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#EEE8E2"
                    ),
                ),
                (
                    "BACKGROUND",
                    (2, 0),
                    (2, -1),
                    colors.HexColor(
                        "#EEE8E2"
                    ),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor(
                        "#D8D1CA"
                    ),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(
        summary_table
    )

    story.append(
        Paragraph(
            "Alerty",
            section_style,
        )
    )

    if alerts:
        rows = [
            [
                "Kategoria",
                "Obiekt",
                "Typ",
                "Termin",
                "Pozostało",
                "Spółka",
                "Oddział",
            ]
        ]

        for alert in alerts:
            rows.append(
                [
                    get_category_name(
                        alert[
                            "category"
                        ]
                    ),
                    Paragraph(
                        (
                            f"<b>{alert['name']}</b>"
                            f"<br/>{alert['secondary']}"
                        ),
                        normal_style,
                    ),
                    Paragraph(
                        alert["type"],
                        normal_style,
                    ),
                    format_date(
                        alert["date"]
                    ),
                    Paragraph(
                        alert[
                            "days_text"
                        ],
                        normal_style,
                    ),
                    Paragraph(
                        (
                            alert.get(
                                "company_name"
                            )
                            or "-"
                        ),
                        normal_style,
                    ),
                    Paragraph(
                        (
                            alert.get(
                                "branch"
                            )
                            or "-"
                        ),
                        normal_style,
                    ),
                ]
            )

        table = Table(
            rows,
            repeatRows=1,
            colWidths=[
                24 * mm,
                48 * mm,
                48 * mm,
                27 * mm,
                38 * mm,
                46 * mm,
                36 * mm,
            ],
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, -1),
                        "DejaVuSans",
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "DejaVuSans-Bold",
                    ),
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#EEE8E2"
                        ),
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.4,
                        colors.HexColor(
                            "#D8D1CA"
                        ),
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        story.append(
            table
        )

    else:
        story.append(
            Paragraph(
                (
                    "Brak alertów spełniających "
                    "wybrane kryteria."
                ),
                normal_style,
            )
        )

    story.append(
        Spacer(
            1,
            8 * mm,
        )
    )

    generated_at = timezone.localtime(
        timezone.now()
    )

    story.append(
        Paragraph(
            (
                "Raport wygenerowano "
                "w systemie Flotivo: "
                f"{generated_at.strftime('%d.%m.%Y %H:%M')}"
            ),
            normal_style,
        )
    )

    document.build(
        story
    )

    output.seek(0)

    return output